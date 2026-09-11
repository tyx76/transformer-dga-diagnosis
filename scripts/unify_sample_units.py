# -*- coding: utf-8 -*-
"""统一样本量纲：全部换算为 μL/L，并修正 sguys99 的键名解析。

输出：
  data/样本/dga_samples.csv            原始量纲（修正解析）
  data/样本/dga_samples_uL_per_L.csv   统一量纲 μL/L
规则：sguys99 的 log10 值 0 表示“该气体为 0/未测”，换算为 0；其余 10^x。
"""
import csv, math, sys
from pathlib import Path
from collections import Counter
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "样本"
RAW = SRC / "dga_samples.csv"
UNI = SRC / "dga_samples_uL_per_L.csv"

HEADER = ["source","h2","ch4","c2h6","c2h4","c2h2","co","co2",
          "scale","source_scale","label_raw","label_std"]

LABEL_MAP = {"正常":"正常","低温过热":"低温过热","中温过热":"中温过热","高温过热":"高温过热",
             "局部放电":"局部放电","低能放电":"低能放电","高能放电":"高能放电"}

def norm_key(k):
    s = str(k).upper()
    for ch in "'\\{}()_ ":
        s = s.replace(ch, "")
    return s

FEATURE_BY_KEY = {
    "LOG10H2":"h2", "LOG10CH4":"ch4", "LOG10C2H2":"c2h2", "LOG10C2H4":"c2h4",
    "LOG10C2H6":"c2h6", "LOG10CO":"co", "LOG10CO2":"co2", "LABELS":"label_raw",
}

def read_alan():
    out = []
    for fn, src in [("alan456__data.xlsx","alan456/data.xlsx"),
                    ("alan456__dataset_589.xlsx","alan456/dataset_(589).xlsx")]:
        p = SRC / fn
        if not p.exists(): continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        data = list(ws.iter_rows(values_only=True))
        wb.close()
        head = [str(c).strip() for c in data[0]]
        idx = {h:i for i,h in enumerate(head)}
        for r in data[1:]:
            if r is None or all(v is None for v in r): continue
            def g(name):
                i = idx.get(name)
                return "" if i is None or r[i] is None else r[i]
            raw = str(g("故障类型")).strip()
            out.append(dict(source=src, h2=g("H2"), ch4=g("CH4"), c2h6=g("C2H6"),
                            c2h4=g("C2H4"), c2h2=g("C2H2"), co="", co2="",
                            source_scale="μL/L", label_raw=raw,
                            label_std=LABEL_MAP.get(raw, raw)))
    return out

def read_sguys99():
    out = []
    for fn, src in [("sguys99__Duval_Classification_1.xlsx","sguys99/Duval_Classification_1.xlsx"),
                    ("sguys99__Duval_Classification_41.xlsx","sguys99/Duval_Classification_41.xlsx")]:
        p = SRC / fn
        if not p.exists(): continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        data = list(ws.iter_rows(values_only=True))
        wb.close()
        if not data: continue
        by_feat = {}
        for r in data:
            feat = FEATURE_BY_KEY.get(norm_key(r[0]))
            if feat: by_feat[feat] = r
        for c in range(1, len(data[0])):
            rec = dict(source=src, source_scale="log10(μL/L)", label_std="")
            for feat in ("h2","ch4","c2h6","c2h4","c2h2","co","co2","label_raw"):
                r = by_feat.get(feat)
                rec[feat] = r[c] if (r is not None and c < len(r)) else ""
            out.append(rec)
    return out

def to_uL(v):
    if v is None or v == "": return ""
    try: f = float(v)
    except (TypeError, ValueError): return ""
    if f == 0: return 0          # log10=0 -> 0/未测
    return round(10 ** f, 6)

rows_raw, rows_uni = [], []
for rec in read_alan():
    rec.update(scale="μL/L")
    rows_uni.append(dict(rec))
    rows_raw.append(dict(rec))
for rec in read_sguys99():
    raw = dict(rec); raw.update(scale="log10(μL/L)")
    rows_raw.append(raw)
    uni = dict(rec); uni.update(scale="μL/L")
    for k in ("h2","ch4","c2h6","c2h4","c2h2","co","co2"):
        uni[k] = to_uL(rec[k])
    rows_uni.append(uni)

def write(path, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(HEADER)
        for r in rows:
            w.writerow([r.get(k, "") for k in HEADER])

write(RAW, rows_raw)
write(UNI, rows_uni)

print("原始量纲表:", RAW, len(rows_raw), "行")
print("统一量纲表:", UNI, len(rows_uni), "行")
print("按来源:", dict(Counter(r["source"] for r in rows_uni)))
sg = [r for r in rows_uni if r["source"].startswith("sguys99")]
for col in ("h2","ch4","c2h6","c2h4","c2h2","co","co2"):
    vals = [float(r[col]) for r in sg if r[col] not in ("", None)]
    if vals:
        print(f"  sguys99 {col}: 非空 {len(vals)}/{len(sg)}, min={min(vals):.4g}, max={max(vals):.4g}")
print("样例(sguys99 首行):", {k: sg[0][k] for k in ("h2","ch4","c2h6","c2h4","c2h2","co","co2","label_raw")})