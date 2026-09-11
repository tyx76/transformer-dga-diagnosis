# -*- coding: utf-8 -*-
"""把 data/样本/ 下的多个公开 DGA 数据集统一成一份 dga_samples.csv。

统一字段：source, h2, ch4, c2h6, c2h4, c2h2, co, co2, scale, label_raw, label_std
说明：
- alan456 数据集为原始浓度(μL/L) + 中文故障类型；
- sguys99 为转置格式(log10(μL/L) 特征 + 数字 Labels)，数字标签含义待确认，label_std 留空；
- sguys99 中 log10 值为 0 表示“该气体为 0/未测”，不还原为浓度。
"""
import csv, sys
from pathlib import Path
from collections import Counter

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import openpyxl

root = Path(__file__).resolve().parent.parent
srcdir = root / "data" / "样本"
out = srcdir / "dga_samples.csv"

HEADER = ["source","h2","ch4","c2h6","c2h4","c2h2","co","co2","scale","label_raw","label_std"]
rows = []

# --- alan456：原始浓度 + 中文标签 ---
LABEL_MAP = {"正常":"正常","低温过热":"低温过热","中温过热":"中温过热","高温过热":"高温过热",
             "局部放电":"局部放电","低能放电":"低能放电","高能放电":"高能放电"}
for fn, src in [("alan456__data.xlsx","alan456/data.xlsx"),
                ("alan456__dataset_589.xlsx","alan456/dataset_(589).xlsx")]:
    p = srcdir / fn
    if not p.exists(): continue
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    data = list(ws.iter_rows(values_only=True))
    head = [str(c).strip() for c in data[0]]
    idx = {h:i for i,h in enumerate(head)}
    for r in data[1:]:
        if r is None or all(v is None for v in r): continue
        def g(name):
            i = idx.get(name)
            return "" if i is None or r[i] is None else r[i]
        raw = str(g("故障类型")).strip()
        rows.append([src, g("H2"), g("CH4"), g("C2H6"), g("C2H4"), g("C2H2"), "", "",
                     "μL/L", raw, LABEL_MAP.get(raw, raw)])
    wb.close()

# --- sguys99：转置格式（行=特征，列=样本）---
FEATS = ["H2","CH4","C2H2","C2H4","C2H6","CO","CO2"]
for fn, src in [("sguys99__Duval_Classification_1.xlsx","sguys99/Duval_Classification_1.xlsx"),
                ("sguys99__Duval_Classification_41.xlsx","sguys99/Duval_Classification_41.xlsx")]:
    p = srcdir / fn
    if not p.exists(): continue
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    data = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(data) < 8: continue
    feat_rows = {}
    for r in data:
        key = str(r[0]).replace("'","")
        for f in FEATS + ["Labels"]:
            if f in key:
                feat_rows[f] = r
    ncols = len(data[0])
    for c in range(1, ncols):
        def v(f):
            r = feat_rows.get(f)
            return "" if r is None or c >= len(r) or r[c] is None else r[c]
        rows.append([src, v("H2"), v("CH4"), v("C2H6"), v("C2H4"), v("C2H2"), v("CO"), v("CO2"),
                     "log10(μL/L)", v("Labels"), ""])

with out.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(HEADER)
    w.writerows(rows)

print("统一数据集:", out)
print("总样本数:", len(rows))
print("按来源:", dict(Counter(r[0] for r in rows)))
print("按 scale:", dict(Counter(r[8] for r in rows)))
print("中文标签分布:", dict(Counter(r[10] for r in rows if r[0].startswith("alan456"))))