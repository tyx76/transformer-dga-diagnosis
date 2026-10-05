#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Run the plant_kb question set through the real RAG/API pipeline."""

from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from vector_kb.citation_verifier import (
    build_retry_prompt,
    remove_invalid_citation_sentences,
    verify_citations,
)
from vector_kb.generation import generate
from vector_kb.retrieval_router import route_and_retrieve

DEFAULT_CASES = ROOT / "data" / "evaluation" / "plant_kb_test_questions.jsonl"
DEFAULT_XLSX = ROOT / "docs" / "04_评测与回归" / "exam_proof" / "plant_kb_api_test_results_20260928.xlsx"
DEFAULT_JSONL = ROOT / "docs" / "04_评测与回归" / "exam_proof" / "plant_kb_api_test_results_20260928.jsonl"


def load_cases(path: Path, limit: int = 0) -> list[dict]:
    rows = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for line in stream:
            if line.strip():
                rows.append(json.loads(line))
    return rows[:limit] if limit > 0 else rows


def _clauses(items: list[dict]) -> str:
    return "；".join(str(item.get("clause") or "-") for item in (items or [])[:5]) or "-"


def run_case(case: dict, top_k: int) -> dict:
    question = str(case.get("question") or "").strip()
    trace: dict[str, object] = {}
    started = time.perf_counter()
    error = ""
    chunks = []
    answer = ""
    verification = {"valid": False, "valid_citations": [], "invalid_citations": []}
    attempts = 0
    route = {}
    try:
        routed = route_and_retrieve(
            question,
            top_k=top_k,
            trace=lambda step, payload: trace.__setitem__(step, payload),
        )
        route = routed.get("route") or {}
        chunks = routed.get("chunks") or []
        if not chunks:
            answer = "资料未覆盖，无法回答"
            verification = {"valid": True, "valid_citations": [], "invalid_citations": []}
        else:
            prompt = question
            for attempt in range(3):
                attempts += 1
                answer = generate(prompt, chunks)
                verification = verify_citations(answer, chunks)
                if verification.get("valid"):
                    break
                if attempt < 2:
                    prompt = build_retry_prompt(question, verification, chunks)
            if not verification.get("valid"):
                answer = remove_invalid_citation_sentences(answer, chunks) or "资料未覆盖，无法回答"
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"

    return {
        "id": case.get("id"),
        "module": case.get("module"),
        "source_question_id": case.get("source_question_id"),
        "question": question,
        "expected_answer": case.get("expected_answer"),
        "source_file": case.get("source_file"),
        "evidence_location": case.get("evidence_location"),
        "intent": trace.get("意图识别") or {},
        "route": route,
        "chunks": chunks,
        "top_clauses": _clauses(chunks),
        "answer": answer,
        "answer_length": len(answer),
        "citation_valid": verification.get("valid"),
        "valid_citations": verification.get("valid_citations") or [],
        "invalid_citations": verification.get("invalid_citations") or [],
        "generation_attempts": attempts,
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "error": error,
    }


def write_outputs(rows: list[dict], xlsx_path: Path, jsonl_path: Path) -> None:
    jsonl_path.parent.mkdir(parents=True, exist_ok=True)
    with jsonl_path.open("w", encoding="utf-8") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False) + "\n")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试结果"
    headers = [
        "编号", "模块", "原始题号", "问题", "意图", "路由 domains", "Top-5条号",
        "引用校验", "有效引用", "无效引用", "生成次数", "耗时(ms)", "错误", "最终回答",
        "期望答案",
    ]
    ws.append(headers)
    for row in rows:
        intent = row.get("intent") or {}
        route = row.get("route") or {}
        ws.append([
            row.get("id"), row.get("module"), row.get("source_question_id"), row.get("question"),
            intent.get("intent"), "、".join(route.get("domains") or []), row.get("top_clauses"),
            "通过" if row.get("citation_valid") else "未通过", "、".join(row.get("valid_citations") or []),
            "、".join(row.get("invalid_citations") or []), row.get("generation_attempts"),
            row.get("elapsed_ms"), row.get("error"), row.get("answer"), row.get("expected_answer"),
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in col) + 3, 70)
        ws.column_dimensions[get_column_letter(col[0].column)].width = width
        for cell in col:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="plant_kb 真实API测试")
    parser.add_argument("--cases", default=str(DEFAULT_CASES))
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    parser.add_argument("--jsonl", default=str(DEFAULT_JSONL))
    parser.add_argument("--top-k", type=int, default=5)
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args(argv)

    cases = load_cases(Path(args.cases), limit=args.limit)
    rows = []
    for index, case in enumerate(cases, start=1):
        row = run_case(case, top_k=args.top_k)
        rows.append(row)
        print(
            f"[{index}/{len(cases)}] {row['id']} {row['module']} "
            f"intent={row['intent'].get('intent')} chunks={len(row['chunks'])} "
            f"citation={'OK' if row['citation_valid'] else 'FAIL'} "
            f"ms={row['elapsed_ms']} error={row['error'] or '-'}",
            flush=True,
        )

    xlsx_path = Path(args.xlsx)
    jsonl_path = Path(args.jsonl)
    if not xlsx_path.is_absolute():
        xlsx_path = ROOT / xlsx_path
    if not jsonl_path.is_absolute():
        jsonl_path = ROOT / jsonl_path
    write_outputs(rows, xlsx_path, jsonl_path)
    print(f"Excel: {xlsx_path}")
    print(f"JSONL: {jsonl_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())