#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""项目回归测试并生成 Excel 报告。"""

from __future__ import annotations

import contextlib
import io
import json
import pickle
import sqlite3
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import main as app_main
from vector_kb.bm25_retriever import INDEX_VERSION, bm25_retrieve
from vector_kb.citation_verifier import remove_invalid_citation_sentences, verify_citations
from vector_kb.domain_guard import is_in_domain
from vector_kb.hybrid_retriever import hybrid_retrieve
from vector_kb.retrieval import retrieve
from vector_kb.rrf_fusion import rrf_fusion

OUTPUT = ROOT / "docs" / "exam_proof" / "回归测试结果_20260918_修订版.xlsx"
CORPUS = ROOT / "data" / "corpus" / "clauses.jsonl"
INDEX = ROOT / "vector_kb" / "bm25_index.pkl"
DB = ROOT / "vector_kb" / "knowledge.db"
RESULTS, RETRIEVAL, CITATIONS = [], [], []


def add_result(case_id, category, name, input_text, expected, actual, passed, ms=0.0, notes=""):
    status = "PASS" if passed is True else "FAIL" if passed is False else "KNOWN"
    RESULTS.append({"id": case_id, "category": category, "name": name, "input": input_text,
                    "expected": expected, "actual": actual, "status": status,
                    "ms": round(ms, 2), "notes": notes})


def clauses(hits):
    return [str(hit.get("clause") or "") for hit in hits]


def fmt_hits(hits):
    parts = []
    for hit in hits:
        score = hit.get("rrf_score", hit.get("score", 0))
        parts.append(f"{hit.get('clause')}({float(score):.4f})")
    return "；".join(parts) or "无"


def retrieval_case(case_id, engine, question, required_any, required_all=None, top_k=5):
    started = time.perf_counter()
    try:
        if engine == "vector":
            hits = retrieve(question, top_k=top_k)
        elif engine == "bm25":
            hits = bm25_retrieve(question, top_k=top_k)
        else:
            hits = hybrid_retrieve(question, top_k=top_k)
        found = set(clauses(hits))
        passed = bool(found & required_any) if required_any else not hits
        if required_all:
            passed = required_all.issubset(found)
        actual = fmt_hits(hits)
        expected = "任一命中：" + "、".join(sorted(required_any))
        if required_all:
            expected += "；全部命中：" + "、".join(sorted(required_all))
        RETRIEVAL.append({"id": case_id, "engine": {"vector": "纯向量", "bm25": "BM25", "hybrid": "混合检索"}[engine],
                          "question": question, "top1": clauses(hits)[0] if hits else "",
                          "top2": clauses(hits)[1] if len(hits) > 1 else "", "top3": clauses(hits)[2] if len(hits) > 2 else "",
                          "top4": clauses(hits)[3] if len(hits) > 3 else "", "top5": clauses(hits)[4] if len(hits) > 4 else "",
                          "scores": [round(float(h.get("rrf_score", h.get("score", 0))), 6) for h in hits],
                          "status": "PASS" if passed else "FAIL"})
    except Exception as exc:
        passed, actual, expected = False, f"{type(exc).__name__}: {exc}", "检索成功并命中目标条文"
    add_result(case_id, "检索", f"{engine} - {question}", question, expected, actual, passed,
               (time.perf_counter() - started) * 1000)


def deterministic_tests():
    valid = ["乙炔超标怎么处理", "C₂H₂注意值", "油里有不好的东西", "变压器油温过高怎么处理"]
    invalid = ["今天晚上吃什么", "变压器怎么炒菜", "周末去看电影"]
    started = time.perf_counter()
    ok = all(is_in_domain(q) for q in valid) and all(not is_in_domain(q) for q in invalid)
    add_result("SYS-001", "领域过滤", "有效与无关问题分类", f"有效={valid}；无关={invalid}",
               "有效 True，无关 False", f"有效={[is_in_domain(q) for q in valid]}；无关={[is_in_domain(q) for q in invalid]}",
               ok, (time.perf_counter() - started) * 1000)

    started = time.perf_counter()
    vec = [{"doc_id": "DOC", "clause": "A", "title": "A", "text": "A", "page": 1, "score": 0.9},
           {"doc_id": "DOC", "clause": "B", "title": "B", "text": "B", "page": 2, "score": 0.8},
           {"doc_id": "DOC", "clause": "C", "title": "C", "text": "C", "page": 3, "score": 0.7}]
    bm = [{"doc_id": "DOC", "clause": "B", "title": "B2", "text": "B", "page": 2, "score": 8.0},
          {"doc_id": "DOC", "clause": "A", "title": "A2", "text": "A", "page": 1, "score": 7.0},
          {"doc_id": "DOC", "clause": "D", "title": "D", "text": "D", "page": 4, "score": 6.0}]
    fused = rrf_fusion(vec, bm, top_k=3)
    ok = clauses(fused) == ["A", "B", "C"] and fused[0]["title"] == "A" and rrf_fusion([], []) == []
    add_result("RRF-001", "RRF", "重复条号累加、排序与空输入", "向量 A/B/C + BM25 B/A/D",
               "A/B 累加并排前列，双空返回 []", f"排序={clauses(fused)}", ok,
               (time.perf_counter() - started) * 1000)

    chunks = [{"doc_id": "DL/T 722-2014", "clause": "9.3.1-表3"},
              {"doc_id": "DL/T 722-2014", "clause": "9.3.2-表4"}]
    cases = [("CIT-001", "真实条号", "结论【依据：DL/T 722-2014 第9.3.1-表3】", True),
             ("CIT-002", "伪造条号", "结论【依据：DL/T 722-2014 第11.5条】", False),
             ("CIT-003", "伪造文档号", "结论【依据：DL/T 999-2020 第9.3.1-表3】", False)]
    for cid, name, answer, expected in cases:
        started = time.perf_counter()
        result = verify_citations(answer, chunks)
        passed = result["valid"] is expected
        CITATIONS.append({"id": cid, "scenario": name, "answer": answer, "valid": result["valid"],
                          "valid_citations": "、".join(result["valid_citations"]),
                          "invalid_citations": "、".join(result["invalid_citations"]),
                          "status": "PASS" if passed else "FAIL"})
        add_result(cid, "引用校验", name, answer, f"valid={expected}",
                   json.dumps(result, ensure_ascii=False), passed, (time.perf_counter() - started) * 1000)

    mixed = "有效句【依据：DL/T 722-2014 第9.3.1-表3】。伪造句【依据：DL/T 722-2014 第11.5条】。"
    cleaned = remove_invalid_citation_sentences(mixed, chunks)
    add_result("CIT-004", "引用校验", "删除无效引用句", mixed, "保留有效句，删除伪造句", cleaned,
               "有效句" in cleaned and "伪造句" not in cleaned)

def ask_tests():
    chunks = [{"doc_id": "DL/T 722-2014", "clause": "9.3.1-表3", "title": "表3", "text": "注意值", "page": 8, "rrf_score": 0.03},
              {"doc_id": "DL/T 722-2014", "clause": "9.3.2-表4", "title": "表4", "text": "产气速率", "page": 9, "rrf_score": 0.02}]
    valid = "处理结论【依据：DL/T 722-2014 第9.3.1-表3】。"
    invalid = "处理结论【依据：DL/T 722-2014 第11.5条】。"

    def fake_hybrid(question, top_k=5, trace=None):
        if trace:
            trace("向量检索Top-K", chunks)
            trace("BM25检索Top-K", chunks[:1])
            trace("RRF融合Top-K", chunks)
        return chunks

    started = time.perf_counter()
    with patch.object(app_main, "hybrid_retrieve", side_effect=fake_hybrid), patch.object(app_main, "generate", return_value=valid):
        output = app_main.ask("正常问题")
    add_result("ASK-001", "主链路桩测试", "首次引用正确直接输出", "正常问题", valid, output,
               output == valid, (time.perf_counter() - started) * 1000)

    calls, responses = [], iter([invalid, valid])
    def retry_generate(prompt, rag_chunks):
        calls.append(prompt)
        return next(responses)
    started = time.perf_counter()
    with patch.object(app_main, "hybrid_retrieve", side_effect=fake_hybrid), patch.object(app_main, "generate", side_effect=retry_generate):
        output = app_main.ask("重写问题")
    passed = output == valid and len(calls) == 2 and "第11.5条" in calls[1]
    add_result("ASK-002", "主链路桩测试", "伪造后一次重写成功", "重写问题",
               "第二次输出正确，反馈包含无效条号", f"调用次数={len(calls)}；输出={output}", passed,
               (time.perf_counter() - started) * 1000)

    started = time.perf_counter()
    with patch.object(app_main, "hybrid_retrieve", side_effect=fake_hybrid), patch.object(app_main, "generate", return_value=invalid):
        output = app_main.ask("持续伪造问题")
    add_result("ASK-003", "主链路桩测试", "两次重写失败后拒答", "持续伪造问题",
               "资料未覆盖，无法回答", output, output == "资料未覆盖，无法回答",
               (time.perf_counter() - started) * 1000)

    buffer = io.StringIO()
    with patch.object(app_main, "hybrid_retrieve", side_effect=fake_hybrid), patch.object(app_main, "generate", return_value=valid):
        with contextlib.redirect_stdout(buffer):
            app_main.ask("调试问题", debug=False)
    debug_buffer = io.StringIO()
    with patch.object(app_main, "hybrid_retrieve", side_effect=fake_hybrid), patch.object(app_main, "generate", return_value=valid):
        with contextlib.redirect_stdout(debug_buffer):
            app_main.ask("调试问题", debug=True)
    debug_text = debug_buffer.getvalue()
    steps = ["用户问题", "查询改写结果", "向量检索Top-K", "BM25检索Top-K", "RRF融合Top-K", "生成调用", "引用校验结果", "最终输出"]
    passed = buffer.getvalue() == "" and all(f"[{step}]" in debug_text for step in steps)
    add_result("DBG-001", "调试输出", "默认静默，开启后输出完整步骤", "调试问题",
               "默认无输出；debug 包含8类步骤", f"默认长度={len(buffer.getvalue())}", passed)


def data_tests():
    started = time.perf_counter()
    rows = [json.loads(line) for line in CORPUS.read_text(encoding="utf-8").splitlines() if line.strip()]
    counts = Counter(row.get("doc_id") for row in rows)
    passed = len(rows) == 123 and counts["DL/T 572-2021"] == 118 and counts["DL/T 722-2014"] == 5
    add_result("DATA-001", "数据一致性", "统一语料条数与文档分布", str(CORPUS),
               "总计123；572=118；722=5", f"总计={len(rows)}；{dict(counts)}", passed,
               (time.perf_counter() - started) * 1000)

    started = time.perf_counter()
    with INDEX.open("rb") as stream:
        payload = pickle.load(stream)
    docs = payload.get("documents") or []
    passed = payload.get("version") == INDEX_VERSION and len(docs) == 123
    add_result("DATA-002", "数据一致性", "BM25索引版本与条数", str(INDEX),
               f"version={INDEX_VERSION}；documents=123", f"version={payload.get('version')}；documents={len(docs)}",
               passed, (time.perf_counter() - started) * 1000)

    started = time.perf_counter()
    con = sqlite3.connect(str(DB))
    count = con.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    con.close()
    add_result("DATA-003", "数据一致性", "向量库块数", str(DB), "123", str(count), count == 123,
               (time.perf_counter() - started) * 1000)

def known_issue_tests():
    hits = hybrid_retrieve("乙炔超标怎么处理", top_k=5)
    dl572 = [hit["clause"] for hit in hits if hit.get("doc_id") == "DL/T 572-2021"]
    add_result("KNOWN-001", "已知问题", "统一语料后的跨文档域干扰", "乙炔超标怎么处理",
               "主命中仍为 DL/T 722", f"混合Top-5={clauses(hits)}；混入572={dl572}", None,
               notes="目标条文已命中，但DL/T 572运维条款仍进入候选，需后续意图路由和文档域过滤。")



def git_value(*args):
    try:
        return subprocess.check_output(["git", *args], cwd=ROOT, text=True, encoding="utf-8").strip()
    except Exception:
        return "unknown"

def build_report(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hf, hfont = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)
    fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"), "FAIL": PatternFill("solid", fgColor="FFC7CE"),
             "KNOWN": PatternFill("solid", fgColor="FFEB9C")}

    def style(ws):
        for cell in ws[1]:
            cell.fill, cell.font = hf, hfont
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
        for col in ws.columns:
            width = max(10, min(max(len(str(c.value or "")) for c in col) + 3, 60))
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
            for cell in col:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    total = len(RESULTS)
    passed = sum(r["status"] == "PASS" for r in RESULTS)
    failed = sum(r["status"] == "FAIL" for r in RESULTS)
    known = sum(r["status"] == "KNOWN" for r in RESULTS)
    summary = wb.create_sheet("汇总")
    summary_rows = [["项目", "结果"], ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                    ["Git提交", git_value("rev-parse", "--short", "HEAD")], ["Python", sys.version.split()[0]],
                    ["openpyxl", openpyxl.__version__], ["用例总数", total], ["通过", passed], ["失败", failed],
                    ["已知问题", known], ["通过率（Pass/(Pass+Fail)）", f"{passed/max(1, passed+failed):.2%}"],
                    ["总体结论", "PASS" if failed == 0 else "FAIL"],
                    ["说明", "检索调用真实Ollama+BM25；生成环节使用桩，避免DeepSeek费用。"]]
    for row in summary_rows:
        summary.append(row)
    style(summary)

    detail = wb.create_sheet("回归明细")
    detail.append(["编号", "类别", "用例", "输入", "期望", "实际", "状态", "耗时(ms)", "备注"])
    for r in RESULTS:
        detail.append([r["id"], r["category"], r["name"], r["input"], r["expected"], r["actual"], r["status"], r["ms"], r["notes"]])
    style(detail)
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 7).fill = fills[detail.cell(row, 7).value]
    ret = wb.create_sheet("检索明细")
    ret.append(["编号", "模式", "问题", "Top1", "Top2", "Top3", "Top4", "Top5", "分数", "状态"])
    for r in RETRIEVAL:
        ret.append([r["id"], r["engine"], r["question"], r["top1"], r["top2"], r["top3"], r["top4"], r["top5"],
                    json.dumps(r["scores"], ensure_ascii=False), r["status"]])
    style(ret)
    for row in range(2, ret.max_row + 1):
        ret.cell(row, 10).fill = fills[ret.cell(row, 10).value]

    cit = wb.create_sheet("引用校验")
    cit.append(["编号", "场景", "回答", "valid", "有效引用", "无效引用", "状态"])
    for r in CITATIONS:
        cit.append([r["id"], r["scenario"], r["answer"], r["valid"], r["valid_citations"], r["invalid_citations"], r["status"]])
    style(cit)
    for row in range(2, cit.max_row + 1):
        cit.cell(row, 7).fill = fills[cit.cell(row, 7).value]

    env = wb.create_sheet("环境")
    env.append(["项目", "值"])
    for row in [["项目根目录", str(ROOT)], ["统一语料", str(CORPUS)], ["BM25索引", str(INDEX)], ["向量库", str(DB)],
                ["向量模型", "Ollama bge-m3:latest"], ["BM25", f"jieba + rank-bm25；INDEX_VERSION={INDEX_VERSION}"],
                ["生成模型", "DeepSeek deepseek-chat（本回归仅使用桩）"]]:
        env.append(row)
    style(env)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    deterministic_tests()
    data_tests()
    ask_tests()

    retrieval_case("VEC-001", "vector", "C₂H₂注意值", {"9.3.1-表3"}, top_k=5)
    retrieval_case("VEC-002", "vector", "变压器油温过高怎么处理", {"7.1.5", "7.1.6", "7.1.7", "7.1.8"}, top_k=5)
    retrieval_case("BM25-001", "bm25", "变压器油温过高怎么处理", {"7.1.5", "7.1.6", "7.1.8"}, top_k=3)
    retrieval_case("BM25-002", "bm25", "乙炔超标怎么处理", {"9.3.1-表3"}, top_k=3)
    retrieval_case("HYB-001", "hybrid", "乙炔超标怎么处理", {"9.3.1-表3"}, top_k=5)
    retrieval_case("HYB-002", "hybrid", "C₂H₂注意值", {"9.3.1-表3"}, top_k=5)
    retrieval_case("HYB-003", "hybrid", "三比值故障类型判断", {"10.2.1-表6", "10.2.1-表7"}, top_k=5)
    retrieval_case("HYB-004", "hybrid", "局部放电", {"10.2.1-表7"}, top_k=5)
    retrieval_case("HYB-005", "hybrid", "变压器油温过高怎么处理", {"7.1.5", "7.1.6", "7.1.7", "7.1.8"},
                   required_all={"7.1.5", "7.1.6", "7.1.8"}, top_k=5)
    retrieval_case("HYB-006", "hybrid", "今天晚上吃什么", set(), top_k=5)

    known_issue_tests()
    build_report(OUTPUT)
    failed = [r for r in RESULTS if r["status"] == "FAIL"]
    print(json.dumps({"output": str(OUTPUT), "total": len(RESULTS),
                      "pass": sum(r["status"] == "PASS" for r in RESULTS), "fail": len(failed),
                      "known": sum(r["status"] == "KNOWN" for r in RESULTS),
                      "failed_cases": [r["id"] for r in failed]}, ensure_ascii=False))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())