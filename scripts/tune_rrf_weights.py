#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Search RRF weights for hybrid-retrieval vs pure-knowledge-base fusion.

The default mode is offline and deterministic: it uses rule-based intent
classification and a citation-structure proxy.  Add ``--allow-llm`` to use the
configured LLM fallback for intent classification and ``--with-generation`` to
measure citation correctness on real generated answers.
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - dependency check at runtime
    raise SystemExit("需要 openpyxl，请执行：pip install -r requirements.txt") from exc

from vector_kb.citation_verifier import verify_citations
from vector_kb.generation import _cite, _load_env, generate
from vector_kb.hybrid_retriever import hybrid_retrieve
from vector_kb.intent_classifier import classify_intent, rule_classify
from vector_kb.intent_router import route_intent
from vector_kb.knowledge_base_adapter import merge_with_hybrid, search_knowledge_base

DEFAULT_CASES = ROOT / "test_cases" / "evaluation_set.jsonl"
FALLBACK_CASES = ROOT / "data" / "evaluation" / "acceptance_cases.jsonl"
DEFAULT_OUTPUT = ROOT / "docs" / "exam_proof" / "RRF权重搜索结果_20260922.xlsx"
WEIGHT_SCHEMES = (
    ("A", 1.0, 1.0),
    ("B", 1.0, 1.2),
    ("C", 1.0, 1.5),
    ("D", 1.0, 0.8),
)


def _resolve_cases_path(explicit: str = "") -> Path:
    """Resolve the requested case file and keep compatibility with old path."""
    candidates = []
    if explicit:
        candidates.append(Path(explicit))
    candidates.extend((DEFAULT_CASES, FALLBACK_CASES))
    for candidate in candidates:
        if not candidate.is_absolute():
            candidate = ROOT / candidate
        if candidate.exists():
            return candidate
    checked = "\n".join(str(path) for path in candidates)
    raise FileNotFoundError(f"找不到评测用例文件，已检查：\n{checked}")


def _load_cases(path: Path) -> list[dict]:
    cases = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for line_number, line in enumerate(stream, start=1):
            if not line.strip():
                continue
            row = json.loads(line)
            if not isinstance(row, dict):
                raise ValueError(f"{path}:{line_number} 不是 JSON 对象")
            row.setdefault("id", f"CASE-{line_number:03d}")
            row.setdefault("question", "")
            row.setdefault("expected_docs", [])
            row.setdefault("expected_clauses", [])
            row.setdefault("category", "")
            cases.append(row)
    if not cases:
        raise ValueError(f"评测用例为空：{path}")
    return cases


def _clauses(items: list[dict] | None) -> list[str]:
    return [str(item.get("clause") or "") for item in (items or []) if isinstance(item, dict)]


def _fmt_hits(items: list[dict] | None) -> str:
    parts = []
    for item in items or []:
        clause = item.get("clause") or "-"
        score = item.get("score", item.get("rrf_score"))
        if score is None:
            parts.append(str(clause))
        else:
            try:
                parts.append(f"{clause}({float(score):.4f})")
            except (TypeError, ValueError):
                parts.append(str(clause))
    return "；".join(parts) or "-"


def _citation_proxy(chunks: list[dict]) -> str:
    """Build a deterministic answer containing only citations from chunks."""
    parts = []
    for chunk in chunks or []:
        if not isinstance(chunk, dict):
            continue
        doc_id = str(chunk.get("doc_id") or "").strip()
        clause = str(chunk.get("clause") or "").strip()
        if doc_id and clause:
            parts.append(f"结论【依据：{_cite(doc_id, clause)}】")
    return "".join(parts)


def _evaluate_citation(question: str, chunks: list[dict], with_generation: bool) -> tuple[bool | None, str, str]:
    """Return (citation_valid, mode, note)."""
    if not chunks:
        return True, "空结果", "无候选条文，无引用可校验"
    if with_generation and _load_env():
        try:
            answer = generate(question, chunks)
            verification = verify_citations(answer, chunks)
            return bool(verification["valid"]), "真实生成", ""
        except Exception as exc:
            note = f"真实生成失败，回退模板：{type(exc).__name__}: {exc}"
    else:
        note = "未启用真实生成，使用检索引用结构代理" if not with_generation else "未找到 API Key，使用检索引用结构代理"
    answer = _citation_proxy(chunks)
    verification = verify_citations(answer, chunks)
    return bool(verification["valid"] and verification["valid_citations"]), "引用结构模板", note


def _classify_for_tuning(question: str, allow_llm: bool) -> dict:
    """Use deterministic rule classification unless LLM fallback is allowed."""
    if allow_llm:
        return classify_intent(question)
    result = rule_classify(question)
    if result is not None:
        return result
    return {"intent": "irrelevant", "confidence": 0.0, "source": "rule_fallback"}


def _run_case(case: dict, with_generation: bool, allow_llm: bool, top_k: int) -> dict:
    question = str(case.get("question") or "").strip()
    expected_clauses = [str(item) for item in (case.get("expected_clauses") or [])]
    expected_docs = [str(item) for item in (case.get("expected_docs") or [])]
    is_refusal_case = not expected_docs and not expected_clauses

    started = time.perf_counter()
    error = ""
    try:
        intent_result = _classify_for_tuning(question, allow_llm)
        route = route_intent(question, classification=intent_result)
        if not question or intent_result.get("intent") == "irrelevant" or route.get("mode") == "refuse":
            hybrid_results, kb_results = [], []
        else:
            hybrid_results = hybrid_retrieve(question, top_k=20)
            kb_results = search_knowledge_base(question, route=route, top_k=20)
    except Exception as exc:
        hybrid_results, kb_results = [], []
        error = f"{type(exc).__name__}: {exc}"
    base_ms = (time.perf_counter() - started) * 1000

    per_scheme = []
    for scheme, w_hybrid, w_kb in WEIGHT_SCHEMES:
        merge_started = time.perf_counter()
        try:
            merged = merge_with_hybrid(
                hybrid_results,
                kb_results,
                top_k=top_k,
                w_hybrid=w_hybrid,
                w_kb=w_kb,
            )
            merge_error = ""
        except Exception as exc:
            merged = []
            merge_error = f"{type(exc).__name__}: {exc}"
        merge_ms = (time.perf_counter() - merge_started) * 1000
        found = set(_clauses(merged))
        hit = any(clause in found for clause in expected_clauses) if expected_clauses else None
        refused = len(merged) == 0
        citation_valid, citation_mode, citation_note = _evaluate_citation(
            question, merged, with_generation=with_generation
        )
        per_scheme.append(
            {
                "scheme": scheme,
                "w_hybrid": w_hybrid,
                "w_kb": w_kb,
                "merged": merged,
                "hit": hit,
                "refused": refused,
                "citation_valid": citation_valid,
                "citation_mode": citation_mode,
                "citation_note": citation_note,
                "elapsed_ms": base_ms + merge_ms,
                "error": merge_error or error,
            }
        )

    return {
        "id": str(case.get("id") or ""),
        "category": str(case.get("category") or ""),
        "question": question,
        "expected_docs": expected_docs,
        "expected_clauses": expected_clauses,
        "is_refusal_case": is_refusal_case,
        "intent": intent_result if "intent_result" in locals() else {},
        "route": route if "route" in locals() else {},
        "hybrid_results": hybrid_results,
        "kb_results": kb_results,
        "schemes": per_scheme,
    }


def _rate(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def _ratio_text(value: float | None) -> str:
    return "N/A" if value is None else f"{value:.2%}"


def _build_summary(rows: list[dict]) -> list[dict]:
    summary = []
    for scheme, w_hybrid, w_kb in WEIGHT_SCHEMES:
        scheme_rows = []
        for row in rows:
            for item in row["schemes"]:
                if item["scheme"] == scheme:
                    scheme_rows.append((row, item))
        hit_denominator = sum(1 for _, item in scheme_rows if item["hit"] is not None)
        hit_numerator = sum(1 for _, item in scheme_rows if item["hit"] is True)
        refusal_denominator = sum(1 for row, _ in scheme_rows if row["is_refusal_case"])
        refusal_numerator = sum(
            1 for row, item in scheme_rows if row["is_refusal_case"] and item["refused"]
        )
        citation_denominator = sum(1 for _, item in scheme_rows if item["citation_valid"] is not None)
        citation_numerator = sum(1 for _, item in scheme_rows if item["citation_valid"] is True)
        elapsed = [item["elapsed_ms"] for _, item in scheme_rows if not item["error"]]
        summary.append(
            {
                "scheme": scheme,
                "w_hybrid": w_hybrid,
                "w_kb": w_kb,
                "cases": len(scheme_rows),
                "hit_n": hit_numerator,
                "hit_denominator": hit_denominator,
                "top5_hit_rate": _rate(hit_numerator, hit_denominator),
                "refusal_n": refusal_numerator,
                "refusal_denominator": refusal_denominator,
                "refusal_rate": _rate(refusal_numerator, refusal_denominator),
                "citation_n": citation_numerator,
                "citation_denominator": citation_denominator,
                "citation_rate": _rate(citation_numerator, citation_denominator),
                "avg_ms": sum(elapsed) / len(elapsed) if elapsed else None,
                "citation_mode": "/".join(sorted({item["citation_mode"] for _, item in scheme_rows})) if scheme_rows else "N/A",
            }
        )
    return summary


def _recommended_scheme(summary: list[dict]) -> str:
    """Pick the best quality scheme; prefer the baseline on a quality tie."""
    order = {scheme: index for index, (scheme, _, _) in enumerate(WEIGHT_SCHEMES)}

    def rank(item: dict):
        return (
            item["top5_hit_rate"] if item["top5_hit_rate"] is not None else -1.0,
            item["citation_rate"] if item["citation_rate"] is not None else -1.0,
            item["refusal_rate"] if item["refusal_rate"] is not None else -1.0,
            -order.get(item["scheme"], 999),
        )

    return max(summary, key=rank)["scheme"]


def _write_excel(path: Path, rows: list[dict], summary: list[dict], cases_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    def style(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = min(max(len(str(cell.value or "")) for cell in col) + 3, 70)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
            for cell in col:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    recommended = _recommended_scheme(summary)
    summary_ws = wb.create_sheet("汇总")
    summary_ws.append(
        [
            "方案",
            "w_hybrid",
            "w_kb",
            "用例数",
            "Top-5命中率",
            "引用正确率",
            "拒答率",
            "平均耗时(ms)",
            "是否推荐",
            "证据模式",
            "命中/可评测",
            "拒答/应拒答",
            "引用正确/可评测",
        ]
    )
    for item in summary:
        summary_ws.append(
            [
                item["scheme"],
                item["w_hybrid"],
                item["w_kb"],
                item["cases"],
                _ratio_text(item["top5_hit_rate"]),
                _ratio_text(item["citation_rate"]),
                _ratio_text(item["refusal_rate"]),
                round(item["avg_ms"], 2) if item["avg_ms"] is not None else "N/A",
                "是" if item["scheme"] == recommended else "",
                item["citation_mode"],
                f'{item["hit_n"]}/{item["hit_denominator"]}',
                f'{item["refusal_n"]}/{item["refusal_denominator"]}',
                f'{item["citation_n"]}/{item["citation_denominator"]}',
            ]
        )
    style(summary_ws)

    detail_ws = wb.create_sheet("逐条明细")
    detail_ws.append(
        [
            "用例",
            "类别",
            "问题",
            "方案",
            "w_hybrid",
            "w_kb",
            "期望条号",
            "hybrid Top-5",
            "知识库 Top-5",
            "融合 Top-5",
            "命中",
            "应拒答",
            "是否拒答",
            "引用正确",
            "耗时(ms)",
            "错误",
            "备注",
        ]
    )
    for row in rows:
        for item in row["schemes"]:
            detail_ws.append(
                [
                    row["id"],
                    row["category"],
                    row["question"],
                    item["scheme"],
                    item["w_hybrid"],
                    item["w_kb"],
                    "、".join(row["expected_clauses"]) or "-",
                    _fmt_hits(row["hybrid_results"][:5]),
                    _fmt_hits(row["kb_results"][:5]),
                    _fmt_hits(item["merged"]),
                    "是" if item["hit"] is True else "否" if item["hit"] is False else "-",
                    "是" if row["is_refusal_case"] else "",
                    "是" if item["refused"] else "否",
                    "是" if item["citation_valid"] is True else "否" if item["citation_valid"] is False else "-",
                    round(item["elapsed_ms"], 2),
                    item["error"],
                    item["citation_note"],
                ]
            )
    style(detail_ws)

    info_ws = wb.create_sheet("说明")
    info_ws.append(["项目", "值"])
    info_ws.append(["评测用例文件", str(cases_path)])
    info_ws.append(["评测用例数", len(rows)])
    info_ws.append(["RRF k", 60])
    info_ws.append(["hybrid候选数", 20])
    info_ws.append(["知识库候选数", 20])
    info_ws.append(["最终Top-K", 5])
    info_ws.append(["权重方案", "A=1.0/1.0, B=1.0/1.2, C=1.0/1.5, D=1.0/0.8"])
    info_ws.append(["引用证据模式", rows[0]["schemes"][0]["citation_mode"] if rows else "N/A"])
    info_ws.append(["候选复用", "四组权重复用同一组hybrid/知识库候选，仅重复执行融合，避免重复Embedding费用。"])
    info_ws.append(["提示", "默认不调用LLM；结果用于离线权重比较，不代表线上最终质量。"])
    style(info_ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="RRF权重搜索：hybrid vs pure_kb")
    parser.add_argument("--cases", default="", help="JSONL评测用例路径")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Excel输出路径")
    parser.add_argument("--top-k", type=int, default=5, help="最终融合Top-K")
    parser.add_argument("--with-generation", action="store_true", help="调用DeepSeek计算真实引用正确率")
    parser.add_argument("--allow-llm", action="store_true", help="规则无法分类时调用LLM兜底")
    args = parser.parse_args(argv)

    cases_path = _resolve_cases_path(args.cases)
    cases = _load_cases(cases_path)
    rows = [
        _run_case(case, with_generation=args.with_generation, allow_llm=args.allow_llm, top_k=args.top_k)
        for case in cases
    ]
    summary = _build_summary(rows)
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = ROOT / output_path
    _write_excel(output_path, rows, summary, cases_path)

    print(f"用例文件：{cases_path}")
    print("方案\tw_hybrid\tw_kb\tTop-5命中率\t引用正确率\t拒答率\t平均耗时(ms)")
    for item in summary:
        avg_ms = f"{item['avg_ms']:.2f}" if item["avg_ms"] is not None else "N/A"
        print(
            f"{item['scheme']}\t{item['w_hybrid']}\t{item['w_kb']}\t"
            f"{_ratio_text(item['top5_hit_rate'])}\t{_ratio_text(item['citation_rate'])}\t"
            f"{_ratio_text(item['refusal_rate'])}\t{avg_ms}"
        )
    print(f"推荐方案：{_recommended_scheme(summary)}")
    print(f"Excel：{output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
